"""
Retrieval-augmented knowledge base — AI-3.

What was here before: an uploaded document was written to disk and
immediately marked "Indexed" — the file was never opened.
`/knowledge-base/search` matched on **filename only**. `/knowledge-base/ask`
scored FAQs by bag-of-words overlap and returned the best one when
`best_score >= 1`, so a **single shared word** — "the", had it not been a
stopword; "appointment" certainly — returned that FAQ as an authoritative
answer to a patient on a live call.

In a clinical setting that last one is a safety problem, not a quality
one, and it is the reason this module leads with a threshold rather than
a ranking. An AI that reliably says "I don't have that, let me transfer
you" is worth more than one that confidently improvises.

The pipeline:

    parse -> chunk -> embed (locally) -> store as VARBINARY
    question -> embed -> cosine over the org's chunks -> threshold
             -> synthesise an answer from ONLY the retrieved text

The response contract is unchanged (`{answer, matched_question, source}`),
so `voice-agent/` needed no modification: the agent already speaks back
`answer` and already treats `source == "none"` as "I don't know". New
source values ("document", "faq") flow through untouched.
"""

import logging
import os
from typing import Any, Optional

from sqlalchemy.orm import Session

from . import embedding_service, llm_service, models, usage_service
from .config import settings

logger = logging.getLogger("rag-service")

# Parsers are imported lazily inside _extract_text so that a missing
# optional dependency degrades one file type rather than breaking import
# of the whole app.
SUPPORTED_TEXT_EXTENSIONS = {".txt", ".md", ".csv"}
SUPPORTED_BINARY_EXTENSIONS = {".pdf", ".docx", ".doc"}


# ---------- Parsing ----------

def _extract_text(path: str) -> Optional[str]:
    """
    Plain text from an uploaded document, or None if it cannot be read.

    Returns None rather than raising: a corrupt or password-protected PDF
    should mark that one document "Failed" and leave the rest of the
    knowledge base working.
    """
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in SUPPORTED_TEXT_EXTENSIONS:
            # errors="replace" rather than strict: a clinic's exported CSV
            # in some legacy encoding is still worth indexing imperfectly.
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()

        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            if reader.is_encrypted:
                logger.warning("PDF is encrypted, cannot index: %s", path)
                return None
            return "\n".join((page.extract_text() or "") for page in reader.pages)

        if ext == ".docx":
            import docx

            document = docx.Document(path)
            parts = [p.text for p in document.paragraphs]
            # Clinic handbooks put opening hours and price lists in tables
            # far more often than in prose, so tables are not optional.
            for table in document.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)

        if ext == ".doc":
            # Legacy binary .doc needs a converter this project does not
            # ship. Say so plainly instead of storing an empty document
            # that silently answers nothing.
            logger.warning("Legacy .doc is not supported for indexing: %s", path)
            return None
    except Exception:
        logger.exception("Failed to extract text from %s", path)
        return None

    logger.warning("No parser for extension %r", ext)
    return None


# ---------- Chunking ----------

def chunk_text(text: str) -> list[str]:
    """
    Splits text into overlapping passages on paragraph boundaries.

    Paragraph-aware rather than a blind character slice: cutting mid-
    sentence produces chunks that embed poorly and read badly when quoted
    back to a caller. The overlap exists so an answer that straddles a
    boundary ("Opening hours\\n\\n9am-5pm...") is not lost to both sides.
    """
    text = (text or "").strip()
    if not text:
        return []

    size = max(settings.kb_chunk_chars, 100)
    overlap = min(max(settings.kb_chunk_overlap_chars, 0), size // 2)

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [text]

    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        if len(current) + len(para) + 2 <= size:
            current = f"{current}\n\n{para}" if current else para
            continue
        if current:
            chunks.append(current)
            # Carry the tail forward so context spans the seam.
            current = (current[-overlap:] + "\n\n" + para) if overlap else para
        else:
            current = para
        # A single paragraph longer than the window still has to be split.
        while len(current) > size:
            chunks.append(current[:size])
            current = current[size - overlap:] if overlap else current[size:]
    if current.strip():
        chunks.append(current.strip())

    return [c for c in chunks if c.strip()]


# ---------- Indexing ----------

def index_document(db: Session, document: models.KBDocument) -> bool:
    """
    Parses, chunks, embeds and stores one document. Returns success.

    Sets `KBDocument.status` to Indexed or Failed — previously it was
    flipped to "Indexed" unconditionally without the file ever being
    opened, so the dashboard reported a working knowledge base that
    contained nothing.

    Re-indexing is safe: existing chunks for the document are removed
    first, so this is idempotent rather than accumulating duplicates.
    """
    try:
        if not document.file_path or not os.path.exists(document.file_path):
            logger.warning("Document %s has no file on disk", document.id)
            document.status = "Failed"
            db.commit()
            return False

        text = _extract_text(document.file_path)
        if not text or not text.strip():
            logger.warning("No extractable text in document %s", document.id)
            document.status = "Failed"
            db.commit()
            return False

        chunks = chunk_text(text)
        if not chunks:
            document.status = "Failed"
            db.commit()
            return False

        vectors = embedding_service.embed_texts(chunks)
        if vectors is None:
            # The model is unavailable. Leave the document marked Failed
            # so the clinic can see it needs re-indexing, rather than
            # storing chunks that can never be retrieved.
            logger.error("Embeddings unavailable — cannot index %s", document.id)
            document.status = "Failed"
            db.commit()
            return False

        db.query(models.KBChunk).filter(
            models.KBChunk.document_id == document.id
        ).delete(synchronize_session=False)

        for i, (chunk, vector) in enumerate(zip(chunks, vectors)):
            db.add(
                models.KBChunk(
                    org_id=document.org_id,
                    document_id=document.id,
                    chunk_index=i,
                    text=chunk,
                    embedding=embedding_service.to_bytes(vector),
                    char_count=len(chunk),
                )
            )

        document.status = "Indexed"
        db.commit()
        logger.info("Indexed document %s into %d chunks", document.id, len(chunks))
        return True
    except Exception:
        logger.exception("Indexing failed for document %s", document.id)
        try:
            db.rollback()
            document.status = "Failed"
            db.commit()
        except Exception:
            pass
        return False


def run_index_task(document_id: str, org_id: str) -> None:
    """
    BackgroundTasks entry point.

    Takes ids and re-queries — by the time this runs the request session
    is closed, and a passed-in ORM instance would be detached.
    """
    from .database import SessionLocal

    db = SessionLocal()
    try:
        document = (
            db.query(models.KBDocument)
            .filter(
                models.KBDocument.id == document_id,
                models.KBDocument.org_id == org_id,
            )
            .first()
        )
        if document is not None:
            index_document(db, document)
    except Exception:
        logger.exception("Index task failed for document %s", document_id)
    finally:
        db.close()


def schedule_index(background_tasks, document_id: str, org_id: str) -> None:
    """Queues indexing so an upload responds immediately."""
    background_tasks.add_task(run_index_task, document_id, org_id)


# ---------- Retrieval ----------

def retrieve(db: Session, org_id: str, question: str, top_k: Optional[int] = None):
    """
    Top-k chunks for a question, above the similarity threshold.

    Scoped to one org — this is the tenant boundary for knowledge, and a
    clinic must never answer from another clinic's handbook.
    """
    query_vector = embedding_service.embed_query(question)
    if query_vector is None:
        return []

    rows = (
        db.query(models.KBChunk)
        .filter(
            models.KBChunk.org_id == org_id,
            models.KBChunk.embedding.isnot(None),
        )
        .all()
    )
    if not rows:
        return []

    import numpy as np

    usable, vectors = [], []
    for row in rows:
        vector = embedding_service.from_bytes(row.embedding)
        if vector is not None:
            usable.append(row)
            vectors.append(vector)
    if not usable:
        return []

    scores = embedding_service.cosine_similarity(query_vector, np.vstack(vectors))
    k = top_k or settings.kb_retrieval_top_k
    ranked = sorted(zip(usable, scores), key=lambda p: float(p[1]), reverse=True)

    return [
        (row, float(score))
        for row, score in ranked[:k]
        if float(score) >= settings.kb_min_similarity
    ]


_SYNTHESIS_SYSTEM = """\
You answer questions for callers to a medical clinic, using ONLY the \
reference text provided.

Rules:
- Answer from the reference text and nothing else. Never add facts from \
general knowledge.
- If the reference text does not contain the answer, reply with exactly: \
INSUFFICIENT_CONTEXT
- Be brief: one or two sentences, spoken aloud to someone on the phone.
- Never give medical advice, a diagnosis, or an interpretation of symptoms.
- Plain sentences only. No markdown, no lists, no citations."""

_INSUFFICIENT = "INSUFFICIENT_CONTEXT"


def _synthesise(question: str, passages: list[str]) -> Optional[str]:
    """
    Composes a spoken answer from retrieved passages.

    Returns None when the model is unconfigured or judges the context
    insufficient — the caller then falls back to the "I don't know" path
    rather than reading a raw document chunk down the phone.
    """
    if not llm_service.is_enabled():
        return None

    context = "\n\n---\n\n".join(passages)[: settings.kb_max_context_chars]
    user = f"Reference text:\n{context}\n\nCaller's question: {question}"

    result = llm_service._call_groq_text(_SYNTHESIS_SYSTEM, user)
    if not result:
        return None
    answer = result.strip()
    if _INSUFFICIENT in answer.upper():
        return None
    return answer


# The single "I don't have that" reply. It is spoken to the caller
# verbatim by the agent, so every path that declines to answer — nothing
# cleared the threshold, the synthesiser returned INSUFFICIENT_CONTEXT,
# or the org's plan excludes the knowledge base — returns exactly this.
# A caller must never be told which of those it was, least of all the
# billing one.
NO_ANSWER_TEXT = (
    "I don't have that specific detail on hand, but I can share our clinic's "
    "general information and contact details, or connect you with our staff."
)


def _no_answer() -> dict[str, Any]:
    return {
        "answer": NO_ANSWER_TEXT,
        "matched_question": None,
        "source": "none",
        "confidence": 0.0,
    }


def answer_question(db: Session, org_id: str, question: str) -> dict[str, Any]:
    """
    The full RAG path, returning the shape `/knowledge-base/ask` promises:
    `{answer, matched_question, source}` — unchanged, so the voice agent
    needed no modification.

    Order of preference:
      1. An indexed document passage above the similarity threshold,
         synthesised into a spoken answer.
      2. A semantically matched FAQ.
      3. An explicit "I don't have that" with `source: "none"`, which is
         the signal the agent already treats as "offer a transfer".

    Nothing below the threshold is ever answered from. That is the whole
    safety property.

    HV-4: an org whose plan excludes `knowledge_base_access` gets branch
    3 immediately. It **degrades rather than 403s** because this function
    serves a live phone call — the agent keys on `source == "none"` and
    offers a transfer, which is a graceful outcome, whereas an error mid
    call is not. The staff-facing writes in `routers/knowledge_base.py`
    do 403, which is where a clinic learns about the restriction.
    """
    if not usage_service.has_feature(db, org_id, "knowledge_base_access"):
        logger.info(
            "KB question declined for org=%s: plan excludes knowledge_base_access",
            org_id,
        )
        return _no_answer()

    hits = retrieve(db, org_id, question)

    if hits:
        passages = [row.text for row, _ in hits]
        answer = _synthesise(question, passages)
        if answer:
            best_row, best_score = hits[0]
            document = (
                db.query(models.KBDocument)
                .filter(models.KBDocument.id == best_row.document_id)
                .first()
                if best_row.document_id
                else None
            )
            return {
                "answer": answer,
                "matched_question": document.name if document else None,
                "source": "document",
                "confidence": round(best_score, 3),
            }

    faq = _best_faq(db, org_id, question)
    if faq:
        faq_row, score = faq
        return {
            "answer": faq_row.answer,
            "matched_question": faq_row.question,
            "source": "faq",
            "confidence": round(score, 3),
        }

    return _no_answer()


def _best_faq(db: Session, org_id: str, question: str):
    """
    The closest FAQ *by meaning*, above the same threshold.

    Replaces keyword overlap with `best_score >= 1`, under which one
    shared word returned an FAQ as authoritative. Falls back to nothing
    (rather than to keyword matching) when embeddings are unavailable:
    answering a caller with a bad match is worse than declining.
    """
    faqs = db.query(models.FAQ).filter(models.FAQ.org_id == org_id).all()
    if not faqs:
        return None

    query_vector = embedding_service.embed_query(question)
    if query_vector is None:
        return None

    # FAQs are few per clinic and change rarely, so they are embedded on
    # demand rather than stored. If a clinic ever has thousands, move
    # them into kb_chunks alongside documents.
    vectors = embedding_service.embed_texts(
        [f"{f.question}\n{f.answer}" for f in faqs]
    )
    if vectors is None:
        return None

    scores = embedding_service.cosine_similarity(query_vector, vectors)
    best = int(scores.argmax())
    if float(scores[best]) < settings.kb_min_similarity:
        return None
    return faqs[best], float(scores[best])
